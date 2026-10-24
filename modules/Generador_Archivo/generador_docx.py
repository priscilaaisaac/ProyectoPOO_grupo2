import io
from typing import Optional
import pandas as pd
from modules.Generador_Archivo.generador_archivo_base import GeneradorArchivo

try:
    import docx
    from docx.shared import Inches, Cm
except ModuleNotFoundError:
    docx = None
    Inches = None
    Cm = None

try:
    from PIL import Image, ImageDraw
except ModuleNotFoundError:
    Image = None
    ImageDraw = None


class GeneradorDOCX(GeneradorArchivo):
    """
    Generador de informes ejecutivos en formato Word (DOCX).
    Hereda de GeneradorArchivo desacoplando la estructura de datos del renderizado físico.
    Incrusta gráficos desde buffers de memoria (BytesIO) con tamaño estático controlado.
    """

    def _generar_grafico_torta_pil(self, df_torta: pd.DataFrame) -> Optional[io.BytesIO]:
        """Genera un gráfico de torta en memoria con PIL si no se proveyó buffer pre-renderizado."""
        if Image is None or ImageDraw is None or df_torta is None or df_torta.empty or "Cantidad" not in df_torta.columns:
            return None
        try:
            cantidades = [float(c) for c in df_torta["Cantidad"].tolist()]
            total = sum(cantidades)
            if total <= 0:
                return None
            img = Image.new("RGB", (620, 320), "white")
            draw = ImageDraw.Draw(img)
            colores = [
                "#1976D2", "#388E3C", "#D32F2F", "#F57C00",
                "#7B1FA2", "#0097A7", "#E64A19", "#5D4037",
                "#455A64", "#C2185B"
            ]

            start_angle = 0.0
            bbox = [40, 35, 290, 285]
            for i, c in enumerate(cantidades):
                angle = (c / total) * 360.0
                draw.pieslice(bbox, start=start_angle, end=start_angle + angle, fill=colores[i % len(colores)], outline="white")
                start_angle += angle

            y = 45
            for i, (name, c) in enumerate(zip(df_torta.iloc[:, 0], cantidades)):
                pct = (c / total) * 100.0
                draw.rectangle([320, y, 336, y + 14], fill=colores[i % len(colores)])
                draw.text((345, y), f"{str(name)[:20]}: {pct:.1f}% ({int(c)})", fill="black")
                y += 24
                if y > 290:
                    break

            buffer = io.BytesIO()
            img.save(buffer, format="PNG")
            buffer.seek(0)
            return buffer
        except Exception:
            return None

    def _generar_grafico_barras_pil(self, df_contingencia: pd.DataFrame) -> Optional[io.BytesIO]:
        """Genera un gráfico de barras agrupadas en memoria con PIL si no se proveyó buffer pre-renderizado."""
        if Image is None or ImageDraw is None or df_contingencia is None or df_contingencia.empty:
            return None
        try:
            df_chart = df_contingencia[df_contingencia.iloc[:, 0] != "Total"].copy()
            cols_series = [c for c in df_chart.columns[1:] if c != "Total"]
            if not cols_series:
                return None
            categories = [str(x) for x in df_chart.iloc[:, 0].tolist()]

            img = Image.new("RGB", (660, 330), "white")
            draw = ImageDraw.Draw(img)
            colores = ["#1976D2", "#388E3C", "#D32F2F", "#F57C00", "#7B1FA2", "#0097A7"]

            vals = []
            for c in cols_series:
                vals.append([float(v) for v in df_chart[c]])
            max_val = max(max(v) for v in vals) if vals and any(vals) else 1.0
            if max_val <= 0:
                max_val = 1.0

            chart_x0, chart_y0 = 60, 35
            chart_x1, chart_y1 = 480, 270
            chart_h = chart_y1 - chart_y0

            draw.line([(chart_x0, chart_y1), (chart_x1, chart_y1)], fill="#333333", width=1)
            draw.line([(chart_x0, chart_y0), (chart_x0, chart_y1)], fill="#333333", width=1)

            n_cats = len(categories)
            n_series = len(cols_series)
            cat_w = (chart_x1 - chart_x0) / max(1, n_cats)
            bar_w = min(24, max(4, (cat_w - 14) / max(1, n_series)))

            for i_cat, cat in enumerate(categories):
                cx = chart_x0 + i_cat * cat_w + 8
                for i_s in range(n_series):
                    val = vals[i_s][i_cat]
                    bh = (val / max_val) * (chart_h - 25)
                    bx0 = cx + i_s * bar_w
                    by0 = chart_y1 - bh
                    bx1 = bx0 + bar_w - 2
                    draw.rectangle([bx0, by0, bx1, chart_y1], fill=colores[i_s % len(colores)])
                draw.text((cx, chart_y1 + 6), str(cat)[:10], fill="black")

            leg_x = 495
            leg_y = 45
            for i_s, s_name in enumerate(cols_series):
                draw.rectangle([leg_x, leg_y, leg_x + 12, leg_y + 12], fill=colores[i_s % len(colores)])
                draw.text((leg_x + 18, leg_y), str(s_name)[:14], fill="black")
                leg_y += 22
                if leg_y > 290:
                    break

            buffer = io.BytesIO()
            img.save(buffer, format="PNG")
            buffer.seek(0)
            return buffer
        except Exception:
            return None

    def _obtener_o_generar_buffer_grafico(self, g: dict) -> Optional[io.BytesIO]:
        """
        Obtiene el buffer de memoria del gráfico o lo genera a partir de su DataFrame.
        Garantiza que el stream esté posicionado al inicio (buffer.seek(0)).
        """
        raw_buf = (
            g.get("buffer")
            or g.get("imagen")
            or g.get("image")
            or g.get("bytes")
            or g.get("figura")
        )
        if isinstance(raw_buf, bytes):
            b = io.BytesIO(raw_buf)
            b.seek(0)
            return b
        if isinstance(raw_buf, io.BytesIO):
            raw_buf.seek(0)
            return raw_buf

        tipo = str(g.get("tipo", ""))
        df_tabla = g.get("tabla")
        if isinstance(df_tabla, pd.DataFrame) and not df_tabla.empty:
            if "Torta" in tipo or "Pie" in tipo:
                return self._generar_grafico_torta_pil(df_tabla)
            else:
                return self._generar_grafico_barras_pil(df_tabla)

        return None

    def generar_archivo(self) -> bytes:
        if docx is None:
            raise ModuleNotFoundError("La librería 'python-docx' no está instalada. Instálala con: pip install python-docx")

        documento = docx.Document()
        documento.add_heading(self.TITULO_OFICIAL, 0)
        
        # Metadatos
        documento.add_paragraph(f"Fecha y hora de emisión: {self.fecha_emision}")
        documento.add_paragraph(f"Dataset Activo: {self.nombre_dataset}")
        documento.add_paragraph(f"Generado por: {self.usuario}")
        
        contador_seccion = 1
        
        # Sección 1: Simulaciones (Solo si hay datos)
        if self.simulaciones:
            documento.add_heading(f"{contador_seccion}. Simulaciones Comerciales y Precios Dinámicos", level=1)
            for idx, sim in enumerate(self.simulaciones, 1):
                documento.add_heading(f"Simulación #{idx}: Ruta: {sim.get('ruta')} | Franja: {sim.get('franja')}", level=2)
                p = documento.add_paragraph()
                p.add_run(f"Precio Base: ${sim.get('precio_base', 0):.2f}\n")
                p.add_run(f"Precio Dinámico Sugerido: ${sim.get('precio_sugerido', 0):.2f}").bold = True
            contador_seccion += 1
              
        # Sección 2: Gráficos y Tablas (Solo si hay datos)
        if self.graficos:
            documento.add_heading(f"{contador_seccion}. Visualizaciones y Gráficos Guardados", level=1)
            for idx, g in enumerate(self.graficos, 1):
                titulo = g.get('titulo') or f"Gráfico #{idx}"
                documento.add_heading(f"Gráfico #{idx}: {titulo}", level=2)

                desc = g.get('descripcion')
                if desc:
                    documento.add_paragraph(str(desc))

                # Incrustar imagen del gráfico en memoria asegurando buffer.seek(0) y tamaño estático
                img_buffer = self._obtener_o_generar_buffer_grafico(g)
                if img_buffer is not None:
                    try:
                        img_buffer.seek(0)
                        ancho_imagen = Inches(5.5) if Inches is not None else None
                        if ancho_imagen is not None:
                            documento.add_picture(img_buffer, width=ancho_imagen)
                        else:
                            documento.add_picture(img_buffer)
                    except Exception:
                        pass

                # Tabla de datos asociada
                df_tabla = g.get('tabla')
                if isinstance(df_tabla, pd.DataFrame) and not df_tabla.empty:
                    tabla_docx = documento.add_table(rows=1, cols=len(df_tabla.columns))
                    tabla_docx.style = 'Table Grid'
                    hdr_cells = tabla_docx.rows[0].cells
                    for i, col in enumerate(df_tabla.columns):
                        hdr_cells[i].text = str(col)
                        if hdr_cells[i].paragraphs and hdr_cells[i].paragraphs[0].runs:
                            hdr_cells[i].paragraphs[0].runs[0].bold = True
                    for _, row in df_tabla.iterrows():
                        row_cells = tabla_docx.add_row().cells
                        for i, val in enumerate(row):
                            row_cells[i].text = f"{val:,.2f}" if isinstance(val, float) else str(val)
            contador_seccion += 1
              
        # Sección 3: Tablas Condicionales (Solo si hay datos)
        if self.tablas_condicionales:
            documento.add_heading(f"{contador_seccion}. Tablas de Probabilidades Condicionales", level=1)
            for idx, tc in enumerate(self.tablas_condicionales, 1):
                documento.add_heading(f"Tabla #{idx}: {tc.get('tipo')}", level=2)
                df_tc = tc.get('dataframe')
                if isinstance(df_tc, pd.DataFrame) and not df_tc.empty:
                    tabla_docx = documento.add_table(rows=1, cols=len(df_tc.columns))
                    tabla_docx.style = 'Table Grid'
                    hdr_cells = tabla_docx.rows[0].cells
                    for i, col in enumerate(df_tc.columns):
                        hdr_cells[i].text = str(col)
                        if hdr_cells[i].paragraphs and hdr_cells[i].paragraphs[0].runs:
                            hdr_cells[i].paragraphs[0].runs[0].bold = True
                    for _, row in df_tc.iterrows():
                        row_cells = tabla_docx.add_row().cells
                        for i, val in enumerate(row):
                            row_cells[i].text = f"{val:,.2f}" if isinstance(val, float) else str(val)
            contador_seccion += 1
              
        # Mensaje si todo está vacío
        if contador_seccion == 1:
            documento.add_paragraph("El analista no ha guardado ningún elemento (simulaciones, gráficos o tablas) para este informe.")
              
        # Guardar en memoria y retornar bytes
        buffer = io.BytesIO()
        documento.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
